"""Bounded document extraction and meeting-scoped evidence retrieval."""
from __future__ import annotations

import asyncio
import hashlib
import io
import multiprocessing
import re
import zipfile
from dataclasses import dataclass
from pathlib import PurePosixPath
from uuid import uuid4

from fastapi import HTTPException, UploadFile
from pydantic import BaseModel

from .repository import _now_iso

MAX_FILE = 10 * 1024 * 1024
MAX_BATCH = 20 * 1024 * 1024
MAX_MEETING_CHARS = 1_000_000


class KnowledgeChunk(BaseModel):
    id: str
    source_id: str
    source: str
    title: str
    filename: str | None = None
    text: str
    locator: str

    @property
    def content(self):
        return self.text


def extract_document(filename: str, data: bytes) -> list[tuple[str, str]]:
    extension = filename.lower().rsplit('.', 1)[-1]
    if extension in {'txt', 'md'}:
        text = data.decode('utf-8-sig')
        return [('lines 1+', text)]
    if extension == 'pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError('Password-protected PDFs are unsupported')
        if len(reader.pages) > 300:
            raise ValueError('PDF exceeds 300 pages')
        return [(f'page {i+1}', page.extract_text() or '') for i, page in enumerate(reader.pages)]
    if extension == 'docx':
        from docx import Document
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(e.file_size for e in entries) > 50 * 1024 * 1024:
                raise ValueError('DOCX expanded content exceeds the limit')
        doc = Document(io.BytesIO(data))
        segments = [(f'paragraph {i+1}', p.text) for i, p in enumerate(doc.paragraphs)]
        for i, table in enumerate(doc.tables):
            for j, row in enumerate(table.rows):
                segments.append((f'table {i+1}, row {j+1}', ' | '.join(c.text for c in row.cells)))
        return segments
    raise ValueError('Use TXT, Markdown, text PDF or DOCX files')


def _extract_worker(connection, filename, data):
    try:
        segments = extract_document(filename, data)
        if sum(len(t) for _, t in segments) > 500_000:
            raise ValueError('Extracted text exceeds 500,000 characters')
        connection.send((True, segments))
    except Exception as error:
        message = str(error) if isinstance(error, (ValueError, UnicodeError)) else 'Cannot extract this document'
        connection.send((False, message))
    finally:
        connection.close()


def bounded_extract(filename: str, data: bytes, timeout: float = 10):
    # Spawned processes can be terminated; a timed-out executor thread cannot.
    context = multiprocessing.get_context('spawn')
    parent, child = context.Pipe(duplex=False)
    process = context.Process(target=_extract_worker, args=(child, filename, data), daemon=True)
    process.start()
    child.close()
    try:
        if not parent.poll(timeout):
            raise ValueError('Document extraction timed out; use a smaller file')
        ok, result = parent.recv()
        if not ok:
            raise ValueError(result)
        return result
    finally:
        parent.close()
        if process.is_alive():
            process.terminate()
        process.join(timeout=1)
        if process.is_alive():
            process.kill()
            process.join()


def chunk_segments(source_id, source, title, filename, segments):
    chunks = []
    for locator, text in segments:
        text = text.replace('\r\n', '\n').replace('\r', '\n').strip()
        for start in range(0, len(text), 1800):
            piece = text[start:start+2000]
            if not piece:
                continue
            chunks.append(KnowledgeChunk(
                id=hashlib.sha256(f'{source_id}:{locator}:{start}'.encode()).hexdigest(),
                source_id=source_id, source=source, title=title, filename=filename,
                text=piece, locator=f'{locator}, characters {start+1}-{start+len(piece)}',
            ))
    return chunks


class KnowledgeService:
    def __init__(self, repository):
        self.repository = repository
        with repository._connect() as db:
            db.executescript('''
                CREATE TABLE IF NOT EXISTS knowledge_sources (
                    id TEXT PRIMARY KEY, session_id TEXT NOT NULL REFERENCES sessions(id),
                    source TEXT NOT NULL, title TEXT NOT NULL, filename TEXT,
                    content_hash TEXT NOT NULL, content TEXT NOT NULL, created_at TEXT NOT NULL,
                    UNIQUE(session_id,content_hash)
                );
                CREATE TABLE IF NOT EXISTS knowledge_chunks (
                    id TEXT PRIMARY KEY, source_id TEXT NOT NULL REFERENCES knowledge_sources(id),
                    text TEXT NOT NULL, locator TEXT NOT NULL
                );
            ''')

    def _store(self, db, session_id, source_id, source, title, filename, segments, created_at):
        content = '\n'.join(t.strip() for _, t in segments if t.strip())
        digest = hashlib.sha256(content.encode()).hexdigest()
        result = db.execute('INSERT OR IGNORE INTO knowledge_sources VALUES(?,?,?,?,?,?,?,?)',
            (source_id, session_id, source, title, filename, digest, content, created_at))
        if not result.rowcount:
            return False
        for chunk in chunk_segments(source_id, source, title, filename, segments):
            db.execute('INSERT INTO knowledge_chunks VALUES(?,?,?,?)', (chunk.id, source_id, chunk.text, chunk.locator))
        return True

    def sync_legacy(self, session_id):
        with self.repository._connect() as db:
            for row in db.execute('SELECT * FROM context_notes WHERE session_id=?', (session_id,)).fetchall():
                if not db.execute('SELECT id FROM knowledge_sources WHERE id=?', (row['id'],)).fetchone():
                    self._store(db, session_id, row['id'], row['source'], row['title'], None,
                        [('note', row['content'])], row['created_at'])

    def list_sources(self, session_id):
        self.sync_legacy(session_id)
        with self.repository._connect() as db:
            rows = db.execute('SELECT * FROM knowledge_sources WHERE session_id=? ORDER BY created_at,id', (session_id,)).fetchall()
        return [dict(id=r['id'], session_id=session_id, source=r['source'], title=r['title'],
            filename=r['filename'], content_hash=r['content_hash'], characters=len(r['content']),
            preview=r['content'][:2000], created_at=r['created_at']) for r in rows]

    def add_notes(self, session_id, notes):
        from .models import ContextNoteView
        self.sync_legacy(session_id)
        created=[]
        with self.repository._connect() as db:
            db.execute('BEGIN IMMEDIATE')
            current=db.execute('SELECT COALESCE(SUM(length(content)),0) FROM knowledge_sources WHERE session_id=?',(session_id,)).fetchone()[0]
            if current+sum(len(n.content) for n in notes)>MAX_MEETING_CHARS:
                raise HTTPException(413,'Meeting knowledge exceeds 1,000,000 characters')
            for note in notes:
                source_id,now=str(uuid4()),_now_iso()
                db.execute('INSERT INTO context_notes VALUES(?,?,?,?,?,?)',(source_id,session_id,note.source,note.title,note.content,now))
                self._store(db,session_id,source_id,note.source,note.title,None,[('note',note.content)],now)
                created.append(ContextNoteView(id=source_id,created_at=now,**note.model_dump()))
        return created

    def search(self, session_id, query, limit=8):
        self.sync_legacy(session_id)
        stop = {'what','when','where','which','who','why','how','does','that','this','with','the','and','our','you','alloy','aloy','about','please','tell','have','your','from','can','are','was','will','its'}
        terms = {t for t in re.findall(r'\w+', query.casefold()) if len(t)>2 and t not in stop}
        with self.repository._connect() as db:
            rows = db.execute('SELECT c.*,s.source,s.title,s.filename FROM knowledge_chunks c JOIN knowledge_sources s ON s.id=c.source_id WHERE s.session_id=?', (session_id,)).fetchall()
        scored = []
        for row in rows:
            words = set(re.findall(r'\w+', (row['title']+' '+row['text']).casefold()))
            score = len(terms & words)
            if score:
                scored.append((score, KnowledgeChunk(**dict(row))))
        scored.sort(key=lambda pair: pair[0], reverse=True)
        return [chunk for _, chunk in scored[:min(limit,8)]]

    def copy_into(self,source_session,target_session):
        """Snapshot full normalized template knowledge, retaining original locators."""
        self.sync_legacy(source_session)
        with self.repository._connect() as db:
            db.execute('BEGIN IMMEDIATE')
            total=db.execute('SELECT COALESCE(SUM(length(content)),0) FROM knowledge_sources WHERE session_id=?',(target_session,)).fetchone()[0]
            for row in db.execute('SELECT * FROM knowledge_sources WHERE session_id=? ORDER BY created_at,id',(source_session,)).fetchall():
                if db.execute('SELECT id FROM knowledge_sources WHERE session_id=? AND content_hash=?',(target_session,row['content_hash'])).fetchone(): continue
                if total+len(row['content'])>MAX_MEETING_CHARS: raise ValueError('Template knowledge exceeds meeting limit')
                source_id=str(uuid4())
                db.execute('INSERT INTO knowledge_sources VALUES(?,?,?,?,?,?,?,?)',(source_id,target_session,row['source'],row['title'],row['filename'],row['content_hash'],row['content'],row['created_at']))
                for chunk in db.execute('SELECT * FROM knowledge_chunks WHERE source_id=?',(row['id'],)).fetchall():
                    chunk_id=hashlib.sha256((source_id+chunk['id']).encode()).hexdigest()
                    db.execute('INSERT INTO knowledge_chunks VALUES(?,?,?,?)',(chunk_id,source_id,chunk['text'],chunk['locator']))
                preview=row['content'] if not row['filename'] and len(row['content'])<=100000 else row['content'][:2000]
                db.execute('INSERT INTO context_notes VALUES(?,?,?,?,?,?)',(source_id,target_session,row['source'],row['title'],preview,row['created_at']))
                total+=len(row['content'])

    async def ingest(self, session_id, text='', files=()):
        if len(files)>5 or len(text)>100_000:
            raise HTTPException(413, 'Use at most five files and 100,000 text characters')
        prepared = []
        if text.strip():
            prepared.append(('briefing', 'Meeting knowledge', None, [('note', text.strip())]))
        total = 0
        for file in files:
            name = PurePosixPath((file.filename or '').replace('\\', '/')).name[:300]
            data = await file.read(MAX_FILE+1)
            total += len(data)
            if len(data)>MAX_FILE or total>MAX_BATCH:
                raise HTTPException(413, 'Limit is 10 MiB per file and 20 MiB combined')
            if name.lower().rsplit('.',1)[-1] not in {'txt','md','pdf','docx'}:
                raise HTTPException(422, f'{name}: use TXT, Markdown, text PDF or DOCX')
            try:
                segments = await asyncio.to_thread(bounded_extract, name, data)
            except Exception as error:
                detail = str(error) if isinstance(error, ValueError) else 'Cannot extract this document'
                raise HTTPException(422, f'{name}: {detail}') from None
            if not any(t.strip() for _, t in segments):
                raise HTTPException(422, f'{name}: no readable text; scanned files need OCR first')
            prepared.append(('file', name, name, segments))
        if not prepared:
            raise HTTPException(422, 'Add text or at least one file')
        self.sync_legacy(session_id)
        duplicates, accepted = [], []
        with self.repository._connect() as db:
            db.execute('BEGIN IMMEDIATE')
            existing = db.execute('SELECT COALESCE(SUM(length(content)),0) FROM knowledge_sources WHERE session_id=?',(session_id,)).fetchone()[0]
            for source,title,filename,segments in prepared:
                content = '\n'.join(t.strip() for _,t in segments if t.strip())
                digest = hashlib.sha256(content.encode()).hexdigest()
                if db.execute('SELECT id FROM knowledge_sources WHERE session_id=? AND content_hash=?',(session_id,digest)).fetchone():
                    duplicates.append(title)
                    continue
                if existing+len(content)>MAX_MEETING_CHARS:
                    raise HTTPException(413, 'Meeting knowledge exceeds 1,000,000 characters')
                source_id, now = str(uuid4()), _now_iso()
                self._store(db,session_id,source_id,source,title,filename,segments,now)
                db.execute('INSERT INTO context_notes VALUES(?,?,?,?,?,?)',(source_id,session_id,source,title,content[:2000],now))
                accepted.append(source_id)
                existing += len(content)
        self.repository.record_event(session_id,'knowledge.added',{'source_ids':accepted,'duplicates':duplicates})
        return {'sources':[s for s in self.list_sources(session_id) if s['id'] in accepted], 'duplicates':duplicates}
